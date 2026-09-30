import openpyxl
p = r"C:\Users\Administrator\Downloads\疯狂动物园2026.8.28.xlsx"
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("=== SHEET:", ws.title, "max_row:", ws.max_row, "max_col:", ws.max_column)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    header = rows[0]
    print("HEADER:", header)
    idx = None
    for i, h in enumerate(header):
        if h and ("获取" in str(h) or "方式" in str(h)):
            idx = i
            break
    print("获取方式列索引:", idx)
    if idx is not None:
        vals = []
        for r in rows[1:]:
            if idx < len(r) and r[idx] is not None and str(r[idx]).strip():
                vals.append(str(r[idx]).strip())
        uniq = []
        for v in vals:
            if v not in uniq:
                uniq.append(v)
        print("获取方式 去重值 (全部):")
        for v in uniq:
            print("  -", repr(v))
        print("唯一数:", len(uniq), " 数据行:", len(vals))
