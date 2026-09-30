import openpyxl

SRC = r"C:\Users\Administrator\Downloads\疯狂动物园2026.8.28.xlsx"
DST = r"C:\Users\Administrator\Downloads\疯狂动物园2026.8.28_完善.xlsx"

# 用户定义的映射：来源关键词(小写) -> 标签
MAPPING = [
    ("自选礼包", "付费"),
    ("徽章商城", "徽章"),
    ("濒危动物", "濒危"),
    ("pvp",     "PVP"),
    ("boss动物", "Boss"),
    ("隐藏动物", "隐藏"),
    ("抽奖轮盘", "轮盘"),
]

def tag_of(cell):
    if cell is None:
        return []
    parts = [p.strip() for p in str(cell).split("\n") if p.strip()]
    tags = []
    for p in parts:
        pl = p.lower()
        for kw, tag in MAPPING:
            if kw in pl and tag not in tags:
                tags.append(tag)
                break
    return tags

wb = openpyxl.load_workbook(SRC)  # 保留原格式

total, filled, combo = 0, 0, 0
for ws in wb.worksheets:
    header = [c.value for c in ws[1]]
    col_idx = None
    for i, h in enumerate(header):
        if h is not None and str(h).strip() == "获取方式":
            col_idx = i
            break
    if col_idx is None:
        print(f"[跳过] sheet「{ws.title}」无精确“获取方式”列")
        continue
    tag_col = col_idx + 2  # 1-based：获取方式列右侧紧邻列
    ws.cell(row=1, column=tag_col, value="属性标签")
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=col_idx + 1).value
        tags = tag_of(cell)
        total += 1
        if tags:
            filled += 1
            if len(tags) > 1:
                combo += 1
            ws.cell(row=r, column=tag_col, value="/".join(tags))
        else:
            ws.cell(row=r, column=tag_col, value="")

wb.save(DST)
print(f"SAVED -> {DST}")
print(f"处理数据行: {total}  有标签: {filled}  组合标签行: {combo}")

# 抽样核对组合行
print("\n--- 组合标签抽样核对 ---")
wb2 = openpyxl.load_workbook(DST, data_only=True)
ws = wb2["动物园动物"]
hdr = [c.value for c in ws[1]]
ti = hdr.index("属性标签") + 1
gi = hdr.index("获取方式") + 1
shown = 0
for r in range(2, ws.max_row + 1):
    g = ws.cell(r, gi).value
    t = ws.cell(r, ti).value
    if t and "/" in str(t):
        print(f"获取方式: {repr(g)}  ->  属性标签: {t}")
        shown += 1
        if shown >= 8:
            break
